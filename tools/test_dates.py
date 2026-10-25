"""Every unit test of both courses, from Windmill's spine (what mrshaffer.org shows), as an .xlsx.   python3 tools/test_dates.py <out.xlsx>
Written 25 Oct 2026 for Croix's "Can I get a download of all test dates as they currently stand." It reads spine/spine.json and the
course repositories beside this one (/home/claude/windy-hill-m7, /home/claude/croix18-windy-hill-a7). Recalculate the workbook after
writing it (its Status and Days-until columns are formulas on TODAY()).
"""
import json, re, ast, datetime as dt, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
S = json.load(open("/home/claude/Windmill/spine/spine.json"))
days = S["days"]
D = dt.date.fromisoformat

def names(repo, d):
    src = open(f"/home/claude/{repo}/{d}/build/course.py", encoding="utf-8").read()
    return {int(k): v for k, v in ast.literal_eval(re.search(r"^UNITS\s*=\s*(\{.*?\})\s*$", src, re.S | re.M).group(1)).items()}
N = {"on": names("windy-hill-m7", "m7"), "acc": names("croix18-windy-hill-a7", "a7")}
# A7's units past 6 are not built yet: their names from the book's contents (A7 IXL Skill Plan and Book Contents.md)
for line in open("/home/claude/croix18-windy-hill-a7/a7/reference/A7 IXL Skill Plan and Book Contents.md", encoding="utf-8"):
    m = re.match(r"## Unit (\d+) (.+?)\s+\(", line)
    if m and int(m.group(1)) not in N["acc"]:
        N["acc"][int(m.group(1))] = m.group(2).strip()

COURSE = {"on": ("M7", "2, 4, 5"), "acc": ("A7", "1, 3")}
rows = []
for c in ("on", "acc"):
    tests, review = {}, {}
    for d in sorted(days):
        e = (days[d] or {}).get(c)
        if not e: continue
        t = re.sub(r"\s*\(as run\)$", "", e["title"])
        if e["kind"] == "exam":
            m = re.match(r"(Units? ([\d–]+)) ", t)
            tests.setdefault(m.group(1), {"units": m.group(2), "dates": [], "asrun": "(as run)" in e["title"]})["dates"].append(D(d))
        elif e["kind"] == "review" and (m := re.match(r"(Units? [\d–]+) review day", t)):
            review[m.group(1)] = D(d)
    for label, t in tests.items():
        us = [int(x) for x in t["units"].split("–")]
        us = list(range(us[0], us[-1] + 1))
        title = "; ".join(N[c].get(u, "") for u in us)
        rows.append({"c": c, "label": label, "title": title, "d1": t["dates"][0],
                     "d2": t["dates"][1] if len(t["dates"]) > 1 else None,
                     "review": review.get(label), "asrun": t["asrun"], "n": len(t["dates"]), "units": us})
rows.sort(key=lambda r: (r["d1"], r["c"] != "on"))

def note(r):
    out = []
    if r["asrun"]: out.append("As run.")
    if len(r["units"]) > 1: out.append("One test for both units.")
    if r["c"] == "on" and not r["review"]: out.append("No review day.")
    if r["d1"] == dt.date(2026, 12, 10) and r["c"] == "on": out.append("Last test of Q2; catch-up week Dec 14–18.")
    if r["d1"] == dt.date(2026, 12, 11) and r["c"] == "acc": out.append("Last test of Q2; semester review Dec 14–16, catch-up Dec 17–18.")
    return " ".join(out)

# ---------- styles
F = "Arial"
ink, sub, rule = "1F2A37", "5B6573", "C9D2DC"
hdr_fill = PatternFill("solid", fgColor="22364F")
band = {"on": PatternFill("solid", fgColor="EEF4FA"), "acc": PatternFill("solid", fgColor="F6F1FA")}
thin = Side(style="thin", color=rule)
DATE = 'ddd, mmm d, yyyy'

def title_block(ws, title, line2, line3, width):
    ws["A1"] = title; ws["A1"].font = Font(name=F, size=15, bold=True, color=ink)
    ws["A2"] = line2; ws["A2"].font = Font(name=F, size=10, color=sub)
    ws["A3"] = line3; ws["A3"].font = Font(name=F, size=10, color=sub)
    for r in (2, 3):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
        ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 15; ws.row_dimensions[3].height = 28

def header(ws, row, cols):
    for i, (name, w) in enumerate(cols, 1):
        cell = ws.cell(row, i, name)
        cell.font = Font(name=F, size=10, bold=True, color="FFFFFF"); cell.fill = hdr_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 20

def body(cell, fill=None, fmt=None, bold=False, color=ink, wrap=False):
    cell.font = Font(name=F, size=10, bold=bold, color=color)
    if fill: cell.fill = fill
    if fmt: cell.number_format = fmt
    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=wrap)
    cell.border = Border(bottom=thin)

def test_sheet(ws, which, title, line2, line3):
    both = which is None
    cols = ([("Course", 8), ("Periods", 9)] if both else []) + [
        ("Unit", 11), ("Unit title", 44), ("Test day 1", 19), ("Test day 2", 19), ("Review day", 19),
        ("Status", 11), ("Days until", 10), ("Notes", 46)]
    title_block(ws, title, line2, line3, len(cols))
    H = 5
    header(ws, H, cols)
    r = H + 1
    for x in rows:
        if not both and x["c"] != which: continue
        vals = ([COURSE[x["c"]][0], COURSE[x["c"]][1]] if both else []) + [x["label"], x["title"], x["d1"], x["d2"],
                (x["review"] if x["review"] else ("None" if x["c"] == "on" else "—"))]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(r, i, v)
            body(cell, fill=band[x["c"]] if both else None, fmt=DATE if isinstance(v, dt.date) else None,
                 bold=(i == (3 if both else 1)), wrap=(i == (4 if both else 2)))
        k = len(vals)
        d1, d2 = get_column_letter(k - 2), get_column_letter(k - 1)
        st = ws.cell(r, k + 1, f'=IF(MAX({d1}{r}:{d2}{r})<TODAY(),"Given",IF({d1}{r}>TODAY(),"Upcoming","Today"))')
        body(st, fill=band[x["c"]] if both else None)
        du = ws.cell(r, k + 2, f'=IF({d1}{r}>TODAY(),{d1}{r}-TODAY(),"")')
        body(du, fill=band[x["c"]] if both else None, fmt="0")
        du.alignment = Alignment(horizontal="center", vertical="top")
        nt = ws.cell(r, k + 3, note(x)); body(nt, fill=band[x["c"]] if both else None, color=sub, wrap=True)
        r += 1
    last = r - 1
    ws.freeze_panes = ws.cell(H + 1, 1)
    ws.auto_filter.ref = f"A{H}:{get_column_letter(len(cols))}{last}"
    r += 1
    foot = ws.cell(r, 1, "Status and Days until work from today's date when the file is opened. Tests given before Sep 24, 2026 "
                         "are not in the plan's calendar, so they are not listed. The plan moves when a class day is lost; "
                         "mrshaffer.org always has the current dates.")
    foot.font = Font(name=F, size=9, italic=True, color=sub); foot.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols)); ws.row_dimensions[r].height = 28
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{H}:{H}"
    ws.sheet_view.showGridLines = False
    return last

asof = "Sun, Oct 25, 2026"
wb = Workbook()
ws = wb.active; ws.title = "All Tests"
test_sheet(ws, None, "Test Dates 2026–27",
           f"Grade 7 Math (M7, periods 2, 4, 5) and Grade 7 Accelerated (A7, periods 1, 3), as the plan stands on {asof}.",
           "An M7 test is one paper over two days, with a review day before day 1 unless noted. An A7 test is one day; "
           "A7 has no review days, and its practice test goes home on the last lesson day.")
test_sheet(wb.create_sheet("M7"), "on", "M7 Test Dates 2026–27",
           f"Grade 7 Math, periods 2, 4, 5, as the plan stands on {asof}.",
           "Each test is one paper over two days, with a review day before day 1 unless noted.")
test_sheet(wb.create_sheet("A7"), "acc", "A7 Test Dates 2026–27",
           f"Grade 7 Accelerated, periods 1 and 3, as the plan stands on {asof}.",
           "Each test is one day. There are no review days; the practice test goes home on the last lesson day.")

# ---------- other dates: state testing, no-class days, catch-up days
ws = wb.create_sheet("State Testing and Other Days")
cols = [("Starts", 19), ("Ends", 19), ("Course", 10), ("What", 52), ("Notes", 52)]
title_block(ws, "State Testing and Other Days", f"Both courses, as the plan stands on {asof}.",
            "Days that matter around the tests: state testing (no class), make-up and catch-up days, and the review before PM3.", len(cols))
header(ws, 5, cols)
other = [
    ("2026-10-08", None, "Both", "State test — no class", "Given."),
    ("2026-12-01", None, "Both", "FAST PM2 (ELA) — no class", "Probable; not yet on the school calendar. If it turns out to be a class day, the plan gets that day back."),
    ("2026-12-08", None, "Both", "FAST PM2 (math) — no class", ""),
    ("2026-12-14", "2026-12-18", "M7", "Catch-up week — make-up tests and missing work", "Nothing new and nothing graded. Q2 and the semester end Fri Dec 18."),
    ("2026-12-14", "2026-12-16", "A7", "Semester review (three days)", "Whiteboard review of the semester's units."),
    ("2026-12-17", "2026-12-18", "A7", "Catch-up days — make-up tests and missing work", "Nothing new and nothing graded. Q2 and the semester end Fri Dec 18."),
    ("2027-04-19", "2027-04-30", "A7", "PM3 review (ten days)", "Content ends with the Unit 17 test, Thu Apr 15."),
    ("2027-04-29", "2027-05-14", "M7", "FAST review block", "Content ends with the Unit 13 test, Wed Apr 28."),
    ("2027-05-03", "2027-05-28", "Both", "FAST PM3 window", "The school sets the test date."),
]
r = 6
for a, b, c, what, nt in other:
    vals = [D(a), D(b) if b else None, c, what, nt]
    for i, v in enumerate(vals, 1):
        cell = ws.cell(r, i, v)
        body(cell, fmt=DATE if isinstance(v, dt.date) else None, bold=(i == 4), color=(sub if i == 5 else ink), wrap=(i >= 4))
    r += 1
ws.freeze_panes = "A6"
ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.sheet_view.showGridLines = False
wb.save(OUT)
print("wrote", OUT, "tests:", len(rows))
for x in rows:
    print(COURSE[x["c"]][0], x["label"], x["d1"], x["d2"], x["review"], "|", x["title"], "|", note(x))
