#!/usr/bin/env python3
"""What a master sheet looks up: every code in its Links tab, with what the fallback names.

    python3 drive/codes.py <master sheet .xlsx>      # JSON: [{"code", "folder", "phrase", "folder_search"}]

drive/test_windyhill.js uses it to hold the two halves to each other: every code a workbook looks
up must be in the tab the script writes, and must point at the file the workbook's fallback names."""
import json, re, sys, urllib.parse
from openpyxl import load_workbook
LOOK = re.compile(r'=IFERROR\("(https://drive\.google\.com/file/d/|https://drive\.google\.com/drive/folders/)"&VLOOKUP\("(k[0-9a-f]{12})",Drive!\$A:\$B,2,0\)(?:&"/view")?,"([^"]+)"\)')
out = []
for row in load_workbook(sys.argv[1])["Links"].iter_rows():
    for c in row:
        m = LOOK.fullmatch(c.value) if isinstance(c.value, str) else None
        if m:
            q = urllib.parse.unquote(m[3].split("?q=", 1)[1])
            out.append(dict(code=m[2], folder=m[1].endswith("folders/"), phrase=re.search(r'title:"([^"]*)"', q)[1], folder_search=q.endswith(" type:folder")))
print(json.dumps(out))
